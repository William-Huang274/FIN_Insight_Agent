"""Content checks for the scripted library report's selected source window."""


def quoted_source_window(report, source_reads=()):
    reference, citation = next(iter(report['citations'].items()))
    quoted = report['narrative_markdown'].split('## 原文摘录\n\n', 1)[1].split('\n\n## 使用说明', 1)[0]
    suffix = f' [{reference}]'
    assert quoted.endswith(suffix), 'fixture_quote_must_bind_its_source'
    quoted = quoted[:-len(suffix)]
    originals = [source['text'] for source in citation['sources']]
    originals.extend(row['passage'] for row in source_reads if row.get('passage_id') == reference)
    assert quoted and any(quoted in text for text in originals), 'quote_not_in_bound_source'
    return quoted


def check_export_texts(report, texts, source_reads=()):
    # PDF/Office preserve the selected quote, not necessarily the whole capture.
    def normalized(text):
        return ''.join(text.replace('**', '').split())
    quoted = normalized(quoted_source_window(report, source_reads))
    return {fmt: {'label': '服务' in text and '验证' in text,
                  'original': quoted in normalized(text)}
            for fmt, text in texts.items()}


def original_reads(state):
    import json
    rows = []
    for message in state['values'].get('messages', []):
        if message.get('type') == 'tool' and message.get('name') == 'read_source_document':
            result = json.loads(message['content'])
            if result.get('operation') == 'read':
                rows.extend(result.get('items', []))
    return rows


def powerpoint_body(path):
    from pptx import Presentation
    # Content slides contain heading, body, page number and brand textboxes.
    # Exclude repeated continuation headings/footers, not the quoted body.
    bodies = []
    for slide in Presentation(path).slides:
        boxes = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text]
        assert len(boxes) == 4, 'unexpected_fixture_slide_structure'
        bodies.append(boxes[1])
    return '\n'.join(bodies)
